from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path


@dataclass(frozen=True)
class DocumentExtraction:
    text: str
    status: str
    kind: str


class DocumentTextExtractor:
    """Bounded local text extraction for common office/document formats.

    This module never sends file contents to a remote service and never OCRs images.
    Image-only PDFs therefore return an empty extraction and are handled by metadata
    fallback rather than pretending that their visual content was understood.
    """

    DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".pptx"}

    def __init__(
        self,
        *,
        max_chars: int = 256_000,
        max_pdf_pages: int = 40,
        max_sheet_cells: int = 12_000,
        max_slides: int = 80,
    ):
        self.max_chars = max(1, int(max_chars))
        self.max_pdf_pages = max(1, int(max_pdf_pages))
        self.max_sheet_cells = max(1, int(max_sheet_cells))
        self.max_slides = max(1, int(max_slides))

    @classmethod
    def supports(cls, filename: str, mime_type: str) -> bool:
        suffix = Path(filename).suffix.lower()
        return suffix in cls.DOCUMENT_EXTENSIONS or mime_type in {
            "application/pdf",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        }

    def extract(self, filename: str, content: bytes, mime_type: str) -> DocumentExtraction:
        suffix = Path(filename).suffix.lower()
        if suffix == ".pdf" or mime_type == "application/pdf":
            return self._extract_pdf(content)
        if suffix == ".docx":
            return self._extract_docx(content)
        if suffix == ".xlsx":
            return self._extract_xlsx(content)
        if suffix == ".pptx":
            return self._extract_pptx(content)
        return DocumentExtraction(text="", status="unsupported", kind="unknown")

    def _bounded(self, parts) -> str:
        collected: list[str] = []
        used = 0
        for part in parts:
            value = str(part or "").strip()
            if not value:
                continue
            remaining = self.max_chars - used
            if remaining <= 0:
                break
            value = value[:remaining]
            collected.append(value)
            used += len(value) + 1
        return "\n".join(collected)[: self.max_chars]

    def _extract_pdf(self, content: bytes) -> DocumentExtraction:
        try:
            from pypdf import PdfReader
        except Exception:
            return DocumentExtraction("", "dependency_missing:pypdf", "pdf")

        try:
            reader = PdfReader(io.BytesIO(content))
            parts = []
            for page in reader.pages[: self.max_pdf_pages]:
                parts.append(page.extract_text() or "")
            text = self._bounded(parts)
            status = "ok:pypdf" if text else "empty_or_image_only:pypdf"
            return DocumentExtraction(text, status, "pdf")
        except Exception as exc:
            return DocumentExtraction("", f"failed:pypdf:{exc.__class__.__name__}", "pdf")

    def _extract_docx(self, content: bytes) -> DocumentExtraction:
        try:
            from docx import Document
        except Exception:
            return DocumentExtraction("", "dependency_missing:python-docx", "docx")

        try:
            document = Document(io.BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = self._bounded(parts)
            return DocumentExtraction(
                text,
                "ok:python-docx" if text else "empty:python-docx",
                "docx",
            )
        except Exception as exc:
            return DocumentExtraction("", f"failed:python-docx:{exc.__class__.__name__}", "docx")

    def _extract_xlsx(self, content: bytes) -> DocumentExtraction:
        try:
            from openpyxl import load_workbook
        except Exception:
            return DocumentExtraction("", "dependency_missing:openpyxl", "xlsx")

        try:
            workbook = load_workbook(
                io.BytesIO(content),
                read_only=True,
                data_only=True,
            )
            parts: list[str] = []
            visited = 0
            for sheet in workbook.worksheets:
                parts.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value) for value in row if value is not None]
                    if values:
                        parts.append(" | ".join(values))
                    visited += len(row)
                    if visited >= self.max_sheet_cells:
                        break
                if visited >= self.max_sheet_cells:
                    break
            workbook.close()
            text = self._bounded(parts)
            return DocumentExtraction(
                text,
                "ok:openpyxl" if text else "empty:openpyxl",
                "xlsx",
            )
        except Exception as exc:
            return DocumentExtraction("", f"failed:openpyxl:{exc.__class__.__name__}", "xlsx")

    def _extract_pptx(self, content: bytes) -> DocumentExtraction:
        try:
            from pptx import Presentation
        except Exception:
            return DocumentExtraction("", "dependency_missing:python-pptx", "pptx")

        try:
            presentation = Presentation(io.BytesIO(content))
            parts: list[str] = []
            for index, slide in enumerate(presentation.slides, start=1):
                if index > self.max_slides:
                    break
                parts.append(f"[Slide {index}]")
                for shape in slide.shapes:
                    text = getattr(shape, "text", "")
                    if text:
                        parts.append(text)
            text = self._bounded(parts)
            return DocumentExtraction(
                text,
                "ok:python-pptx" if text else "empty:python-pptx",
                "pptx",
            )
        except Exception as exc:
            return DocumentExtraction("", f"failed:python-pptx:{exc.__class__.__name__}", "pptx")
