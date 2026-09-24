import io
import unittest

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from aegisflow.context.content_classifier import ContentClassifier


class DocumentContentIntelligenceTests(unittest.TestCase):
    def test_docx_financial_content_is_extracted(self):
        document = Document()
        document.add_paragraph(
            "Bank transaction report. Payment invoice and account number details."
        )
        buffer = io.BytesIO()
        document.save(buffer)

        result = ContentClassifier(enable_vision=False).classify(
            filename="report.docx",
            content=buffer.getvalue(),
            mime_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )

        self.assertEqual(result.category, "financial")
        self.assertEqual(result.document_kind, "docx")
        self.assertTrue(result.document_status.startswith("ok:"))
        self.assertIn("local_document_text", result.analysis_mode)

    def test_xlsx_credential_content_is_extracted(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["setting", "value"])
        sheet.append(["api key", "api_key = demo-secret"])
        buffer = io.BytesIO()
        workbook.save(buffer)

        result = ContentClassifier(enable_vision=False).classify(
            filename="settings.xlsx",
            content=buffer.getvalue(),
            mime_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

        self.assertEqual(result.category, "credentials")
        self.assertEqual(result.document_kind, "xlsx")
        self.assertTrue(result.document_status.startswith("ok:"))

    def test_pptx_public_event_content_is_extracted(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        text_box = slide.shapes.add_textbox(0, 0, 5_000_000, 1_000_000)
        text_box.text = "Open Registration Seminar and Workshop"
        buffer = io.BytesIO()
        presentation.save(buffer)

        result = ContentClassifier(enable_vision=False).classify(
            filename="slides.pptx",
            content=buffer.getvalue(),
            mime_type=(
                "application/vnd.openxmlformats-officedocument."
                "presentationml.presentation"
            ),
        )

        self.assertEqual(result.category, "public")
        self.assertEqual(result.document_kind, "pptx")
        self.assertTrue(result.document_status.startswith("ok:"))


if __name__ == "__main__":
    unittest.main()
